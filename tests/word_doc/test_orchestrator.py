from apps.word_doc_services.doc_creation_orchestrator.orchestrator import WordDocumentOrchestrator

from apps.triage.slugs import *
from docx.shared import Pt, Inches
from apps.word_doc_services.common_resources import (
    get_mhclg_green_rgb,
    _bold_font_name
)

general_answers_corporate_spend_exit: dict = {
    total_value_of_business_case: "between-12k-and-2m",
    novel_repercussive_contentious_hmt_consent: "no",
    is_this_request_a_pilot: "no",
    making_a_change_to_or_additional_money_for_existing_business_case: "no",
    where_is_the_budget_held: "Any Answer",
    which_option_describes_what_you_are_trying_to_do: procure_goods_and_services_from_third_party,
    which_best_describes_your_spend: spend_on_corporate_activities,
    give_your_bjc_a_title: "Test BJC Title",
    provide_a_high_level_summary: "details"
}

general_answers_something_else_exit: dict = {
    total_value_of_business_case: "between-12k-and-2m",
    novel_repercussive_contentious_hmt_consent: "no",
    is_this_request_a_pilot: "no",
    making_a_change_to_or_additional_money_for_existing_business_case: "no",
    where_is_the_budget_held: "Any Answer",
    which_option_describes_what_you_are_trying_to_do: procuring_something_else,
    give_your_bjc_a_title: "Test BJC Title",
    provide_a_high_level_summary: "details"
}

general_answers_hrbp_exit_dict: dict = {
    total_value_of_business_case: "between-12k-and-2m",
    novel_repercussive_contentious_hmt_consent: "no",
    is_this_request_a_pilot: "no",
    making_a_change_to_or_additional_money_for_existing_business_case: "no",
    where_is_the_budget_held: "Any Answer",
    which_option_describes_what_you_are_trying_to_do: hire_contracted_workers_to_fill_temporary_capacity_gap,
    hire_contracted_workers_to_fill_temporary_capacity_gap_notice: "yes",
    give_your_bjc_a_title: "Test BJC Title",
    provide_a_high_level_summary: "details"
}

procurement_answers_dict: dict = {
    total_value_of_business_case: "between-12k-and-2m",
    novel_repercussive_contentious_hmt_consent: "no",
    is_this_request_a_pilot: "no",
    making_a_change_to_or_additional_money_for_existing_business_case: "no",
    where_is_the_budget_held: "Any Answer",
    which_option_describes_what_you_are_trying_to_do: procure_goods_and_services_from_third_party,
    which_best_describes_your_spend: "Any Answer",
    give_your_bjc_a_title: "Test BJC Title",
    provide_a_high_level_summary: "details"
}

grant_answers = {
    total_value_of_business_case: "between-12k-and-2m",
    novel_repercussive_contentious_hmt_consent: "no",
    is_this_request_a_pilot: "no",
    making_a_change_to_or_additional_money_for_existing_business_case: "no",
    any_other_business_cases_that_are_connected_to_this_work: "Any Answer",
    where_is_the_budget_held: "Any Answer",
    which_option_describes_what_you_are_trying_to_do: grant_choice, 
    grant_notice: "yes",
    grant_covered_countries: "Any Answer",
    eligible_organisations: "Any Answer",
    grant_routes: formula_grant_choice,
    funding_allocation: "*",
    give_your_bjc_a_title: "Test BJC Title",
    provide_a_high_level_summary: "details"
}


def test_orchestrator_for_procurement_route():
    # Arrange
    procurement_headers: list = [
        "This is your business justification case template for Test BJC Title",
        "1. Participants and organisational details",
        "3. Strategic case: why is this needed?",
        "4. Options analysis: why is this the best option and how does it show value for money?",
        "5. Commercial case: how will you procure this?",
        "6. Financial case - can we afford it?",
        "7. Management case: can this be delivered?",
        "8. Other business cases",
        "9. Subject Matter Expert (SME) assurance recommendations and declaration",
        "Additional SMEs and reviewers (optional):",
        "10. SRO or SCS approval",
        "11. Supporting documents"
    ]

    orc = WordDocumentOrchestrator(procurement_answers_dict)

    # Act
    orc.populate_word_doc()
    doc = orc.wrapper.doc

    # Assert
    # if we made it to the formatting it has completed the creation
    for s in doc.sections:
        assert s.bottom_margin == Inches(1)
        assert s.top_margin == Inches(1)
        assert s.left_margin == Inches(1)
        assert s.right_margin == Inches(1)

    # ensure every header and formatting is complete
    for p in doc.paragraphs:
        if p.text in procurement_headers:
            r = p.runs[0]
            if (
                r.font.size == Pt(16)
                and r.font.color.rgb == get_mhclg_green_rgb()
                and r.font.bold == True
                and r.font.italic == False
                and r.font.name == _bold_font_name
            ):
                procurement_headers.remove(p.text)

    assert len(procurement_headers) == 0


def test_orchestrator_for_grants_route():
    # Arrange
    grants_headers: list = [
            "This is your business justification case template for Test BJC Title",
            "1. Participants and organisational details",
            "3. Economic Case",
            "4. Commercial case: how will you procure this?",
            "5. Financial case - can we afford it?",
            "6. Management case: can this be delivered?",
            "7. Other business cases",
            "8. Subject Matter Expert (SME) assurance recommendations and declaration",
            "Additional SMEs and reviewers (optional):",
            "9. SRO or SCS approval",
            "10. Supporting documents"
        ]

    orc = WordDocumentOrchestrator(grant_answers)
    
    # Act
    orc.populate_word_doc()
    doc = orc.wrapper.doc

    # Assert
    # if we made it to the formatting it has completed the creation
    for s in doc.sections:
        assert s.bottom_margin == Inches(1)
        assert s.top_margin == Inches(1)
        assert s.left_margin == Inches(1)
        assert s.right_margin == Inches(1)

        for p in doc.paragraphs:
                if p.text in grants_headers:
                    r = p.runs[0]
                    if (
                        r.font.size == Pt(16)
                        and r.font.color.rgb == get_mhclg_green_rgb()
                        and r.font.bold == True
                        and r.font.italic == False
                        and r.font.name == _bold_font_name
                    ):
                        grants_headers.remove(p.text)
        
        assert len(grants_headers) == 0


def test_orchestrator_for_general_hrbp_route():
    # Arrange
    procurement_headers: list = [
        "This is your business justification case template for Test BJC Title",
        "1. Participants and organisational details",
        "3. Strategic case: why is this needed?",
        "4. Options analysis: why is this the best option and how does it show value for money?",
        "5. Commercial case: how will you procure this?",
        "6. Financial case - can we afford it?",
        "7. Management case: can this be delivered?",
        "8. Other business cases",
        "9. Subject Matter Expert (SME) assurance recommendations and declaration",
        "Additional SMEs and reviewers (optional):",
        "10. SRO or SCS approval",
        "11. Supporting documents"
    ]
    
    orc = WordDocumentOrchestrator(general_answers_hrbp_exit_dict)

    # Act
    orc.populate_word_doc()
    doc = orc.wrapper.doc

    # Assert
    # if we made it to the formatting it has completed the creation
    for s in doc.sections:
        assert s.bottom_margin == Inches(1)
        assert s.top_margin == Inches(1)
        assert s.left_margin == Inches(1)
        assert s.right_margin == Inches(1)

    # ensure every header and formatting is complete
    for p in doc.paragraphs:
        if p.text in procurement_headers:
            r = p.runs[0]
            if (
                r.font.size == Pt(16)
                and r.font.color.rgb == get_mhclg_green_rgb()
                and r.font.bold == True
                and r.font.italic == False
                and r.font.name == _bold_font_name
            ):
                procurement_headers.remove(p.text)

    assert len(procurement_headers) == 0


def test_orchestrator_for_general_something_else_route():
    # Arrange
    procurement_headers: list = [
        "This is your business justification case template for Test BJC Title",
        "1. Participants and organisational details",
        "3. Strategic case: why is this needed?",
        "4. Options analysis: why is this the best option and how does it show value for money?",
        "5. Commercial case: how will you procure this?",
        "6. Financial case - can we afford it?",
        "7. Management case: can this be delivered?",
        "8. Other business cases",
        "9. Subject Matter Expert (SME) assurance recommendations and declaration",
        "Additional SMEs and reviewers (optional):",
        "10. SRO or SCS approval",
        "11. Supporting documents"
    ]
    
    orc = WordDocumentOrchestrator(general_answers_something_else_exit)

    # Act
    orc.populate_word_doc()
    doc = orc.wrapper.doc

    # Assert
    # if we made it to the formatting it has completed the creation
    for s in doc.sections:
        assert s.bottom_margin == Inches(1)
        assert s.top_margin == Inches(1)
        assert s.left_margin == Inches(1)
        assert s.right_margin == Inches(1)

    # ensure every header and formatting is complete
    for p in doc.paragraphs:
        if p.text in procurement_headers:
            r = p.runs[0]
            if (
                r.font.size == Pt(16)
                and r.font.color.rgb == get_mhclg_green_rgb()
                and r.font.bold == True
                and r.font.italic == False
                and r.font.name == _bold_font_name
            ):
                procurement_headers.remove(p.text)

    assert len(procurement_headers) == 0


def test_orchestrator_for_general_corporate_spend_route():
    # Arrange
    procurement_headers: list = [
        "This is your business justification case template for Test BJC Title",
        "1. Participants and organisational details",
        "3. Strategic case: why is this needed?",
        "4. Options analysis: why is this the best option and how does it show value for money?",
        "5. Commercial case: how will you procure this?",
        "6. Financial case - can we afford it?",
        "7. Management case: can this be delivered?",
        "8. Other business cases",
        "9. Subject Matter Expert (SME) assurance recommendations and declaration",
        "Additional SMEs and reviewers (optional):",
        "10. SRO or SCS approval",
        "11. Supporting documents"
    ]
    
    orc = WordDocumentOrchestrator(general_answers_corporate_spend_exit)

    # Act
    orc.populate_word_doc()
    doc = orc.wrapper.doc

    # Assert
    # if we made it to the formatting it has completed the creation
    for s in doc.sections:
        assert s.bottom_margin == Inches(1)
        assert s.top_margin == Inches(1)
        assert s.left_margin == Inches(1)
        assert s.right_margin == Inches(1)

    # ensure every header and formatting is complete
    for p in doc.paragraphs:
        if p.text in procurement_headers:
            r = p.runs[0]
            if (
                r.font.size == Pt(16)
                and r.font.color.rgb == get_mhclg_green_rgb()
                and r.font.bold == True
                and r.font.italic == False
                and r.font.name == _bold_font_name
            ):
                procurement_headers.remove(p.text)

    assert len(procurement_headers) == 0

